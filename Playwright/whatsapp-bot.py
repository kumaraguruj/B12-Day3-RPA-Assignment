#!/usr/bin/env python3
"""Send WhatsApp messages to contacts in Contact.xlsx.

Run `python whatsapp-bot.py` to preview or `python whatsapp-bot.py --send` to
send. Enter the message when prompted. On first use, scan the WhatsApp Web QR
code; the browser profile is saved for future runs.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


def find_search_box(page):
    # WhatsApp has changed its search markup over time; try accessible names first.
    search_locators = (
        page.get_by_role("textbox", name=re.compile(r"Search", re.I)),
        page.get_by_placeholder(re.compile(r"Search", re.I)),
        page.locator('[contenteditable="true"][aria-label*="Search" i]'),
    )

    def visible_search_box():
        for locator in search_locators:
            for index in range(locator.count()):
                candidate = locator.nth(index)
                try:
                    candidate.wait_for(state="visible", timeout=750)
                    if candidate.is_editable():
                        return candidate
                except PlaywrightTimeoutError:
                    continue
        return None

    search_box = visible_search_box()
    if search_box:
        return search_box

    # The search textbox may only appear after opening the sidebar search UI.
    search_button = page.get_by_role(
        "button", name=re.compile(r"Search|New chat", re.I)
    ).first
    try:
        search_button.click(timeout=1_000)
    except PlaywrightTimeoutError:
        pass

    search_box = visible_search_box()
    if search_box:
        return search_box

    raise PlaywrightTimeoutError("Could not find WhatsApp's contact search box.")


def normalize_contact_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def wait_for_contact_chat(page, name: str) -> None:
    page.wait_for_function(
        """expected => {
            const headers = document.querySelectorAll('header');
            const activeHeader = headers[headers.length - 1];
            const normalize = value => value.toLowerCase().replace(/[^a-z0-9]/g, '');
            return activeHeader && normalize(activeHeader.innerText).includes(expected);
        }""",
        arg=normalize_contact_name(name),
        timeout=5_000,
    )


def search_and_open_contact(page, name: str, phone: str | None) -> None:
    """Search WhatsApp Web by phone, then by name, and open a matching result."""
    normalized_name = normalize_contact_name(name)
    name_pattern = re.compile(
        r"[\W_]*".join(re.escape(char) for char in normalized_name), re.I
    )
    phone_suffix = phone[-7:] if phone else None
    queries = [(phone, True)] if phone else []
    queries.append((name, False))

    for query, is_phone_search in queries:
        search = find_search_box(page)
        try:
            search.fill(query, timeout=1_000)
        except PlaywrightTimeoutError:
            search = find_search_box(page)
            search.fill(query, timeout=1_000)

        # Wait for this contact's result, not a stale row from the previous search.
        rows = page.locator('[role="listitem"]')
        result_pattern = name_pattern
        if is_phone_search and phone_suffix:
            phone_pattern = r"\D*".join(phone_suffix)
            result_pattern = re.compile(
                rf"(?:{name_pattern.pattern}|{phone_pattern})", re.I
            )
        matching_row = rows.filter(has_text=result_pattern).first
        try:
            matching_row.wait_for(state="visible", timeout=3_000)
            matching_row.click()
            wait_for_contact_chat(page, name)
            return
        except PlaywrightTimeoutError:
            pass

        # A phone search can show only the contact name, without the number.
        text_result = page.get_by_text(name_pattern).last
        try:
            text_result.wait_for(state="visible", timeout=1_000)
            text_result.click()
            wait_for_contact_chat(page, name)
            return
        except PlaywrightTimeoutError:
            continue

    number_info = f" (+{phone})" if phone else ""
    raise RuntimeError(f"No WhatsApp search result found for {name}{number_info}.")


def find_message_box(page):
    for locator in (
        page.get_by_role("textbox", name=re.compile(r"Type a message", re.I)).last,
        page.locator('footer [contenteditable="true"]').last,
        page.locator('[contenteditable="true"][data-tab]').last,
    ):
        try:
            locator.wait_for(state="visible", timeout=4_000)
            return locator
        except PlaywrightTimeoutError:
            pass
    raise PlaywrightTimeoutError("Could not find the message box after opening the contact.")


def wait_for_composer_clear(page, composer, timeout: int = 10_000) -> None:
    deadline = time.monotonic() + timeout / 1_000
    while time.monotonic() < deadline:
        try:
            if not composer.inner_text(timeout=500).strip():
                return
        except PlaywrightTimeoutError:
            pass
        page.wait_for_timeout(100)
    raise PlaywrightTimeoutError("Message composer did not clear after sending.")


def read_contacts(path: Path) -> list[tuple[str, str | None]]:
    """Read names from column A and optional phone numbers from column B."""
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    contacts: list[tuple[str, str | None]] = []
    for row in workbook.active.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or not str(row[0]).strip():
            continue
        name = str(row[0]).strip()
        raw_phone = row[1] if len(row) > 1 else None
        if isinstance(raw_phone, (int, float)) and float(raw_phone).is_integer():
            raw_phone = int(raw_phone)
        phone = re.sub(r"\D", "", str(raw_phone)) if raw_phone is not None else ""
        contacts.append((name, phone or None))
    workbook.close()
    if not contacts:
        raise ValueError("No contacts found in the workbook.")
    return contacts


def write_reports(
    app_path: Path,
    workbook_path: Path,
    message: str,
    mode: str,
    contacts: list[dict[str, str | None]],
    run_error: str | None = None,
) -> tuple[Path, Path]:
    created_at = datetime.now().astimezone()
    report_date = created_at.strftime("%Y-%m-%d")
    counts = {
        status: sum(contact["status"] == status for contact in contacts)
        for status in ("sent", "preview", "failed", "pending")
    }
    details = {
        "report_date": report_date,
        "created_at": created_at.isoformat(timespec="seconds"),
        "mode": mode,
        "workbook": str(workbook_path.resolve()),
        "message": message,
        "summary": {"total_contacts": len(contacts), **counts},
        "run_error": run_error,
        "contacts": contacts,
    }

    json_path = app_path / f"whatsapp_report_{report_date}.json"
    json_path.write_text(
        json.dumps(details, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    excel_path = app_path / f"whatsapp_report_{report_date}.xlsx"
    report = Workbook()
    summary = report.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary_rows = [
        ("Report date", report_date),
        ("Created at", details["created_at"]),
        ("Mode", mode),
        ("Workbook", details["workbook"]),
        ("Message", message),
        ("Total contacts", len(contacts)),
        ("Sent", counts["sent"]),
        ("Preview", counts["preview"]),
        ("Failed", counts["failed"]),
        ("Pending", counts["pending"]),
        ("Run error", run_error or ""),
    ]
    for row in summary_rows:
        summary.append(row)

    contact_sheet = report.create_sheet("Contacts")
    contact_sheet.append(["Name", "Phone", "Status", "Message", "Error"])
    for contact in contacts:
        contact_sheet.append(
            [
                contact["name"],
                contact["phone"] or "",
                contact["status"],
                contact["message"],
                contact["error"] or "",
            ]
        )

    for sheet in report.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            sheet.column_dimensions[column[0].column_letter].width = width

    report.save(excel_path)
    print(f"Reports saved: {json_path.name}, {excel_path.name}")
    return json_path, excel_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--contacts", type=Path, default=None,
                        help="Excel workbook (defaults to Contact.xlsx beside this script)")
    parser.add_argument(
        "--message",
        help="Message to send; prompts interactively if omitted",
    )
    parser.add_argument(
        "--send", action="store_true", help="Send messages instead of previewing"
    )
    parser.add_argument(
        "--profile",
        type=Path,
        default=Path.home() / ".whatsapp_playwright_profile",
        help="Persistent Chromium profile directory",
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=2.0,
        help="Seconds to wait between messages (default: 2)",
    )
    args = parser.parse_args()

    # Get the message from the command line or prompt interactively.
    message = args.message
    if message is None:
        if not sys.stdin.isatty():
            parser.error("provide --message when running without an interactive terminal")
        message = input("Enter the message to send: ").strip()
    if not message:
        parser.error("message cannot be empty")

    # Load contacts from the selected workbook or the app directory default.
    workbook_path = args.contacts
    if workbook_path is None:
        app_path = Path(__file__).resolve().parent
        candidates = [app_path / "Contact.xlsx", app_path / "contact.xslx"]
        workbook_path = next((path for path in candidates if path.is_file()), candidates[0])
    contacts = read_contacts(workbook_path)
    report_contacts = [
        {
            "name": name,
            "phone": phone,
            "message": message.replace("{name}", name),
            "status": "pending" if args.send else "preview",
            "error": None,
        }
        for name, phone in contacts
    ]

    # Show the recipients and stop here unless sending was requested.
    print(f"Workbook: {workbook_path}")
    print(f"Message: {message}")
    print(f"Recipients ({len(contacts)}):")
    for contact in report_contacts:
        name = contact["name"]
        phone = contact["phone"]
        target = f"+{phone}" if phone else "name search"
        print(f"  {contact['status'].upper()}: {name} — {target}")
    if not args.send:
        print("\nPreview only. Add --send to send these messages.")
        write_reports(
            Path(__file__).resolve().parent,
            workbook_path,
            message,
            "preview",
            report_contacts,
        )
        return 0

    sent = 0
    run_error = None
    try:
        # Open WhatsApp Web with a persistent browser profile.
        with sync_playwright() as playwright:
            context = playwright.chromium.launch_persistent_context(
                user_data_dir=str(args.profile), headless=False,
            )
            try:
                page = context.pages[0] if context.pages else context.new_page()
                page.goto("https://web.whatsapp.com/", wait_until="domcontentloaded")
                print(
                    "Scan the QR code if needed. Waiting up to 5 minutes for WhatsApp Web login."
                )

                # Wait until the logged-in search control is available.
                login_deadline = time.monotonic() + 300
                while True:
                    try:
                        find_search_box(page)
                        break
                    except PlaywrightTimeoutError:
                        if time.monotonic() >= login_deadline:
                            raise RuntimeError(
                                "WhatsApp Web login timed out. Scan the QR code and try again."
                            )
                        page.wait_for_timeout(1_500)

                # Send each personalized message and pause between contacts.
                for contact in report_contacts:
                    name = contact["name"]
                    phone = contact["phone"]
                    try:
                        search_and_open_contact(page, name, phone)
                        composer = find_message_box(page)
                        composer.fill(contact["message"])
                        composer.press("Enter")
                        wait_for_composer_clear(page, composer)
                        contact["status"] = "sent"
                        sent += 1
                        recipient = f" (+{phone})" if phone else ""
                        print(f"SENT {sent}/{len(contacts)}: {name}{recipient}")
                        if args.delay > 0 and sent < len(contacts):
                            time.sleep(args.delay)
                    except Exception as exc:
                        contact["status"] = "failed"
                        contact["error"] = str(exc)
                        print(f"FAILED: {name} - {exc}", file=sys.stderr)
                        raise
            finally:
                context.close()
    except Exception as exc:
        run_error = f"{type(exc).__name__}: {exc}"
        raise
    finally:
        write_reports(
            Path(__file__).resolve().parent,
            workbook_path,
            message,
            "send",
            report_contacts,
            run_error,
        )

    print(f"Finished: {sent} message(s) sent.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, ValueError, RuntimeError, PlaywrightTimeoutError) as exc:
        raise SystemExit(f"Error: {exc}")
